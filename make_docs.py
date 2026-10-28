"""Generate report figures (assets/figures/*.png) and presentation.pptx.

Every number is computed by game_theory.py, so the slides always match the app.
Run:  .venv/bin/python make_docs.py
"""
from __future__ import annotations

from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from matplotlib.colors import ListedColormap
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

import game_theory as gt
from game_theory import BASELINE, HIGH, LOW, STRATEGIES
from team import FACULTY, PROJECT_TITLE, SECTION, SUBJECT, TEAM

ROOT = Path(__file__).parent
FIG = ROOT / "assets" / "figures"
SHOTS = ROOT / "assets" / "screenshots"
TEAL, AMBER, NAVY, GREY = "#0f7c8a", "#d98a1c", "#16324f", "#b8c4cc"
PROFILES = [(a, b) for a in STRATEGIES for b in STRATEGIES]
AN = gt.analyze(BASELINE)


def rs(v: float) -> str:
    return f"₹{v:,.0f}"


# --------------------------------------------------------------------------
# Figures
# --------------------------------------------------------------------------
def _style(ax):
    ax.spines[["top", "right"]].set_visible(False)
    ax.grid(axis="y", color="#e6ecef")
    ax.set_axisbelow(True)


def _save(fig, name: str, dpi: int = 200) -> Path:
    path = FIG / name
    fig.tight_layout()
    fig.savefig(path, dpi=dpi)
    plt.close(fig)
    return path


def fig_profits() -> Path:
    labels = [gt.profile_label(k) for k in PROFILES]
    a = [AN.matrix[k][0] for k in PROFILES]
    b = [AN.matrix[k][1] for k in PROFILES]
    x = np.arange(4)
    fig, ax = plt.subplots(figsize=(7, 4))
    ba = ax.bar(x - 0.2, a, 0.38, color=TEAL, label="Station A")
    bb = ax.bar(x + 0.2, b, 0.38, color=AMBER, hatch="//", edgecolor="white", label="Station B")
    for bars in (ba, bb):
        ax.bar_label(bars, labels=[rs(v) for v in bars.datavalues], fontsize=9, padding=2)
    ax.set_xticks(x, labels)
    ax.set_ylabel("Profit (₹)")
    ax.set_xlabel("Strategy profile (A / B)")
    ax.set_title("Profit of each station in every outcome (baseline)", loc="left", color=NAVY)
    ax.legend(frameon=False)
    ax.set_ylim(0, max(a + b) * 1.18)
    _style(ax)
    return _save(fig, "fig_profits.png")


def fig_joint() -> Path:
    labels = [gt.profile_label(k) for k in PROFILES]
    vals = [AN.joint[k] for k in PROFILES]
    colors = [NAVY if k in AN.joint_max else GREY for k in PROFILES]
    fig, ax = plt.subplots(figsize=(7, 4))
    bars = ax.bar(labels, vals, color=colors)
    tags = [rs(v) + (" ★ max" if k in AN.joint_max else "") + (" • NE" if k in AN.equilibria else "")
            for k, v in zip(PROFILES, vals)]
    ax.bar_label(bars, labels=tags, fontsize=9, padding=2)
    ax.set_ylabel("Combined profit (₹)")
    ax.set_xlabel("Strategy profile (A / B)")
    ax.set_title("Combined profit: Nash equilibrium vs joint maximum", loc="left", color=NAVY)
    ax.set_ylim(0, max(vals) * 1.18)
    _style(ax)
    return _save(fig, "fig_joint.png")


def fig_q_sensitivity() -> Path:
    qs = np.linspace(0.5, 1.0, 201)
    rows = gt.sensitivity_q(BASELINE, qs)
    vs_low = np.array([r["A Low / Low"] - r["A High / Low"] for r in rows])
    vs_high = np.array([r["A Low / High"] - r["A High / High"] for r in rows])
    lo, hi = gt.q_thresholds(BASELINE)
    fig, ax = plt.subplots(figsize=(8, 4.2))
    ax.axvspan(0.5, lo, color=AMBER, alpha=0.15, label="(High, High) only")
    ax.axvspan(lo, hi, color=NAVY, alpha=0.10, label="Both (Low, Low) and (High, High)")
    ax.axvspan(hi, 1.0, color=TEAL, alpha=0.13, label="(Low, Low) only")
    ax.plot(qs, vs_low, color=TEAL, lw=2.5, label="Gain from Low vs rival Low")
    ax.plot(qs, vs_high, color=NAVY, lw=2.5, ls="--", label="Gain from Low vs rival High")
    ax.axhline(0, color="#555", lw=0.8)
    ax.axvline(BASELINE.cheap_share, color="#333", ls=":", lw=1)
    ax.text(BASELINE.cheap_share + 0.005, vs_low.min(), "baseline q = 0.80", fontsize=9)
    ax.set_xlabel("q = share won by the cheaper station")
    ax.set_ylabel("Profit(Low) − Profit(High), ₹")
    ax.set_title(f"Station A's incentive to undercut (thresholds q = {lo:.3f} and {hi:.3f})",
                 loc="left", color=NAVY)
    ax.legend(fontsize=8, frameon=False, loc="upper left")
    ax.set_xlim(0.5, 1.0)
    _style(ax)
    return _save(fig, "fig_q_sensitivity.png")


def fig_heatmap() -> Path:
    qs = np.linspace(0.5, 1.0, 101)
    costs = np.linspace(0, 9.9, 100)
    codes = gt.sensitivity_grid(BASELINE, qs, costs, "cost")
    present = sorted(set(codes.flatten().tolist()))
    palette = {0: TEAL, 1: AMBER, 2: NAVY, 3: "#c0504d", 4: "#999999"}
    index = {c: i for i, c in enumerate(present)}
    z = np.vectorize(index.get)(codes)
    fig, ax = plt.subplots(figsize=(8, 4.4))
    ax.imshow(z, origin="lower", aspect="auto", extent=[qs[0], qs[-1], costs[0], costs[-1]],
              cmap=ListedColormap([palette[c] for c in present]), interpolation="nearest")
    marker, = ax.plot(BASELINE.cheap_share, BASELINE.cost_a, "wX", ms=11, mec="black")
    handles = [plt.Rectangle((0, 0), 1, 1, color=palette[c]) for c in present]
    ax.legend(handles + [marker], [gt.EQ_CODES[c] for c in present] + ["Baseline"],
              fontsize=8, loc="upper left", framealpha=0.9)
    ax.set_xlabel("q = share won by the cheaper station")
    ax.set_ylabel("Operating cost, both stations (₹/kWh)")
    ax.set_title("Pure Nash equilibria across q and cost (p_L = ₹10, p_H = ₹12)", loc="left", color=NAVY)
    return _save(fig, "fig_heatmap_cost.png")


def fig_matrix() -> Path:
    fig, ax = plt.subplots(figsize=(6.4, 3.6))
    ax.axis("off")
    ax.set_xlim(0, 3); ax.set_ylim(0, 3)
    ax.text(1.5, 2.75, "B: Low (₹10)", ha="center", fontsize=12, weight="bold", color=NAVY)
    ax.text(2.5, 2.75, "B: High (₹12)", ha="center", fontsize=12, weight="bold", color=NAVY)
    for r, sa in enumerate(STRATEGIES):
        y = 1.9 - r
        ax.text(0.95, y, f"A: {sa}\n(₹{BASELINE.price(sa):g})", ha="right", va="center", fontsize=12,
                weight="bold", color=NAVY)
        for c, sb in enumerate(STRATEGIES):
            x = 1.05 + c
            ne = (sa, sb) in AN.equilibria
            ax.add_patch(plt.Rectangle((x, y - 0.45), 0.9, 0.9, fc="#eef5f7" if ne else "white",
                                       ec=NAVY if ne else GREY, lw=3 if ne else 1.2))
            ua, ub = AN.matrix[(sa, sb)]
            a_br, b_br = sa in AN.br_a[sb], sb in AN.br_b[sa]
            ax.text(x + 0.45, y + 0.12, f"A {rs(ua)}{' ▲' if a_br else ''}", ha="center", fontsize=11,
                    color=TEAL, weight="bold" if a_br else "normal")
            ax.text(x + 0.45, y - 0.15, f"B {rs(ub)}{' ◆' if b_br else ''}", ha="center", fontsize=11,
                    color="#9a5c00", weight="bold" if b_br else "normal")
            if ne:
                ax.text(x + 0.45, y - 0.37, "Nash equilibrium", ha="center", fontsize=8, color=NAVY)
            elif (sa, sb) in AN.joint_max:
                ax.text(x + 0.45, y - 0.37, "★ max joint profit", ha="center", fontsize=8, color=NAVY)
    ax.text(0.05, 0.25, "▲ = A's best response   ◆ = B's best response   thick border = Nash equilibrium",
            fontsize=8.5, color="#333")
    return _save(fig, "fig_payoff_matrix.png", dpi=220)


# --------------------------------------------------------------------------
# Presentation
# --------------------------------------------------------------------------
def rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_.lstrip("#"))


W, H = Inches(13.333), Inches(7.5)
FONT = "Calibri"


def _text(slide, x, y, w, h, text, size=18, bold=False, color="#22313f", align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(text if isinstance(text, list) else [text]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.size, run.font.bold, run.font.name = Pt(size), bold, FONT
        run.font.color.rgb = rgb(color)
    return box


def _bullets(slide, x, y, w, h, items, size=18):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = "•  " + item
        run.font.size, run.font.name = Pt(size), FONT
        run.font.color.rgb = rgb("#22313f")
        p.space_after = Pt(10)
    return box


def _rect(slide, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE, line=None):
    r = slide.shapes.add_shape(shape, x, y, w, h)
    r.fill.solid(); r.fill.fore_color.rgb = rgb(fill)
    if line:
        r.line.color.rgb = rgb(line)
    else:
        r.line.fill.background()
    return r


def _frame(prs, kicker: str, title: str, number: int):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _rect(s, 0, 0, Inches(0.18), H, TEAL)
    _text(s, Inches(0.6), Inches(0.35), Inches(11), Inches(0.4), kicker.upper(), 12, True, TEAL)
    _text(s, Inches(0.6), Inches(0.7), Inches(12), Inches(0.9), title, 32, True, NAVY)
    _text(s, Inches(11.8), Inches(7.0), Inches(1.2), Inches(0.3), f"{number} / 10", 11, False, "#7a8b97",
          PP_ALIGN.RIGHT)
    return s


def _notes(slide, text: str):
    slide.notes_slide.notes_text_frame.text = text


def _picture(slide, path: Path, x, y, w=None, h=None):
    if path.exists():
        return slide.shapes.add_picture(str(path), x, y, w, h)
    return _text(slide, x, y, w or Inches(4), Inches(0.5), f"[missing image: {path.name}]", 14)


def _style_cell(cell, fill, color, size, bold):
    cell.fill.solid()
    cell.fill.fore_color.rgb = rgb(fill)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    for p in cell.text_frame.paragraphs:
        p.alignment = PP_ALIGN.CENTER
        for run in p.runs:
            run.font.name, run.font.size, run.font.bold = FONT, Pt(size), bold
            run.font.color.rgb = rgb(color)


def _matrix_table(slide, x, y, w, h):
    t = slide.shapes.add_table(3, 3, x, y, w, h).table
    hdr = ["A \\ B", f"B: Low (₹{BASELINE.low_price:g})", f"B: High (₹{BASELINE.high_price:g})"]
    for c, v in enumerate(hdr):
        t.cell(0, c).text = v
        _style_cell(t.cell(0, c), NAVY, "#ffffff", 16, True)
    for r, sa in enumerate(STRATEGIES, start=1):
        t.cell(r, 0).text = f"A: {sa} (₹{BASELINE.price(sa):g})"
        _style_cell(t.cell(r, 0), NAVY, "#ffffff", 16, True)
        for c, sb in enumerate(STRATEGIES, start=1):
            ua, ub = AN.matrix[(sa, sb)]
            ne = (sa, sb) in AN.equilibria
            tag = "\nNash equilibrium" if ne else ("\n★ max joint profit" if (sa, sb) in AN.joint_max else "")
            t.cell(r, c).text = f"({rs(ua)}, {rs(ub)}){tag}"
            _style_cell(t.cell(r, c), "#d7eef1" if ne else "#ffffff", NAVY, 20, ne)


def build_presentation(figs: dict[str, Path]) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    lo, hi = gt.q_thresholds(BASELINE)
    eq = AN.equilibria[0]
    pd_a = AN.pd.a
    gap = AN.joint[AN.joint_max[0]] - AN.joint[eq]

    # 1 Title
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _rect(s, 0, 0, W, H, NAVY)
    _rect(s, Inches(0.8), Inches(2.3), Inches(1.2), Inches(0.08), AMBER)
    _text(s, Inches(0.8), Inches(1.6), Inches(11), Inches(0.5), f"{SUBJECT.upper()} PROJECT · {SECTION.upper()}",
          14, True, "#9fd3da")
    _text(s, Inches(0.8), Inches(2.6), Inches(11.5), Inches(1.8), PROJECT_TITLE, 40, True, "#ffffff")
    _text(s, Inches(0.8), Inches(4.3), Inches(11), Inches(0.5),
          "A two-player pricing game: payoff matrix, best responses, Nash equilibrium", 18, False, "#d5e6ea")
    _text(s, Inches(0.8), Inches(5.1), Inches(8), Inches(1.8), [f"{n}  ·  {i}" for n, i in TEAM], 16, False,
          "#ffffff")
    _text(s, Inches(8.8), Inches(5.1), Inches(4), Inches(0.5), f"Faculty: {FACULTY}", 16, False, "#d5e6ea")
    _notes(s, "Introduce the team and the topic: we model how two EV charging stations choose prices when "
              "each one's profit depends on the other's choice.")

    # 2 Problem
    s = _frame(prs, "Problem statement", "Why EV charging stations face a pricing trade-off", 2)
    _bullets(s, Inches(0.6), Inches(1.8), Inches(7.2), Inches(4.8), [
        "Two charging stations serve the same drivers and each sets a price per kWh.",
        "A higher price earns more on every kWh sold…",
        "…but drivers move to the cheaper station, so it sells fewer kWh.",
        "Each station's best price therefore depends on the rival's price.",
        "That interdependence is what game theory analyses.",
    ], 20)
    _rect(s, Inches(8.3), Inches(2.0), Inches(4.4), Inches(3.6), "#eef5f7", MSO_SHAPE.ROUNDED_RECTANGLE, TEAL)
    m_lo, m_hi = BASELINE.low_price - BASELINE.cost_a, BASELINE.high_price - BASELINE.cost_a
    _text(s, Inches(8.6), Inches(2.25), Inches(3.9), Inches(3.2), [
        "Margin per kWh", f"Low: ₹{BASELINE.low_price:g} − ₹{BASELINE.cost_a:g} = ₹{m_lo:g}",
        f"High: ₹{BASELINE.high_price:g} − ₹{BASELINE.cost_a:g} = ₹{m_hi:g}", "",
        "Market share", "Equal prices: 50% each",
        f"Cheaper: {BASELINE.cheap_share:.0%} · dearer: {1 - BASELINE.cheap_share:.0%}"], 18, False, NAVY)
    _notes(s, f"Explain the trade-off with the margins: High earns ₹{m_hi:g} per kWh instead of ₹{m_lo:g}, but "
              "if the rival is cheaper you keep only 20% of the market.")

    # 3 Objectives
    s = _frame(prs, "Objectives and scope", "What the model sets out to establish", 3)
    _text(s, Inches(0.6), Inches(1.75), Inches(12), Inches(0.9),
          "Research question: which prices will competing stations choose, and how does that compare with "
          "the prices that maximise their combined profit?", 20, True, TEAL)
    _bullets(s, Inches(0.6), Inches(2.9), Inches(12), Inches(4), [
        "Formulate the competition as a 2-player, 2-strategy simultaneous game.",
        "Derive every payoff from a demand-allocation and profit model, in code.",
        "Compute best responses and all pure-strategy Nash equilibria algorithmically.",
        "Compare the equilibrium with joint-profit maximisation; test the Prisoner's Dilemma conditions.",
        "Re-solve the game across parameter ranges and present it in an interactive website.",
        "Scope: an illustrative model with assumed prices and costs, not an empirical forecast.",
    ], 19)
    _notes(s, "State the research question clearly. Emphasise that the numbers are illustrative assumptions.")

    # 4 Formulation
    s = _frame(prs, "Game-theory formulation", "Players, strategies and game type", 4)
    rows = [("Players", "Station A, Station B"),
            ("Strategies", f"Low (₹{BASELINE.low_price:g}/kWh) or High (₹{BASELINE.high_price:g}/kWh)"),
            ("Payoff", "Profit = (price − cost) × energy sold"),
            ("Timing", "Simultaneous: each chooses without seeing the other"),
            ("Type", "Non-cooperative, finite, normal-form, complete information")]
    tbl = s.shapes.add_table(len(rows), 2, Inches(0.6), Inches(1.9), Inches(12), Inches(4.2)).table
    tbl.columns[0].width, tbl.columns[1].width = Inches(2.8), Inches(9.2)
    for r, (k, v) in enumerate(rows):
        for c, val in enumerate((k, v)):
            tbl.cell(r, c).text = val
            _style_cell(tbl.cell(r, c), NAVY if c == 0 else ("#f3f7f8" if r % 2 else "#ffffff"),
                        "#ffffff" if c == 0 else "#22313f", 20, c == 0)
            tbl.cell(r, c).text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
    _notes(s, "Walk through each row. Non-cooperative means no binding agreements; simultaneous means "
              "neither station observes the other's price before choosing.")

    # 5 Math model
    s = _frame(prs, "Mathematical model", "Demand allocation and profit equations", 5)
    _rect(s, Inches(0.6), Inches(1.8), Inches(7.3), Inches(3.9), "#f3f7f8", MSO_SHAPE.ROUNDED_RECTANGLE)
    _text(s, Inches(0.9), Inches(2.0), Inches(6.9), Inches(3.6), [
        "s_A = 0.5     if p_A = p_B",
        "s_A = q       if p_A < p_B",
        "s_A = 1 − q   if p_A > p_B", "",
        "E_A = D · s_A        E_B = D · (1 − s_A)",
        "π_A = (p_A − c_A) · E_A",
        "π_B = (p_B − c_B) · E_B",
    ], 22, False, NAVY)
    _bullets(s, Inches(8.3), Inches(1.9), Inches(4.6), Inches(4.8), [
        f"D = {BASELINE.demand:,.0f} kWh (fixed)",
        f"p_L = ₹{BASELINE.low_price:g}, p_H = ₹{BASELINE.high_price:g}",
        f"c_A = c_B = ₹{BASELINE.cost_a:g}/kWh",
        f"q = {BASELINE.cheap_share:.0%} to the cheaper station",
        "No fixed costs, capacity or queues",
    ], 18)
    o = AN.outcomes[(LOW, HIGH)]
    _notes(s, f"Worked example, Low/High: A gets {o.a.share:.0%} of {BASELINE.demand:,.0f} = {o.a.energy:,.0f} kWh "
              f"and earns ({o.a.price:g}−{BASELINE.cost_a:g})×{o.a.energy:,.0f} = {rs(o.a.profit)}; B gets "
              f"{o.b.energy:,.0f} kWh and earns {rs(o.b.profit)}.")

    # 6 Payoff matrix
    s = _frame(prs, "Payoff matrix", "Four outcomes, computed from the formulas", 6)
    _matrix_table(s, Inches(0.6), Inches(1.9), Inches(8.2), Inches(3.9))
    _bullets(s, Inches(9.1), Inches(1.9), Inches(3.9), Inches(4.5), [
        f"If B plays Low, A's best reply is {' or '.join(AN.br_a[LOW])} "
        f"({rs(AN.matrix[(LOW, LOW)][0])} vs {rs(AN.matrix[(HIGH, LOW)][0])}).",
        f"If B plays High, A's best reply is {' or '.join(AN.br_a[HIGH])} "
        f"({rs(AN.matrix[(LOW, HIGH)][0])} vs {rs(AN.matrix[(HIGH, HIGH)][0])}).",
        "The game is symmetric, so B's best replies mirror A's.",
    ], 17)
    _text(s, Inches(0.6), Inches(6.0), Inches(8), Inches(0.5), "Cells show (Profit A, Profit B).", 14, False,
          "#5a6b78")
    _notes(s, "Read one cell aloud. Then show that whatever B does, Low gives A more profit: a dominant strategy.")

    # 7 Nash
    s = _frame(prs, "Nash equilibrium", f"Why ({eq[0]}, {eq[1]}) is the unique equilibrium", 7)
    _bullets(s, Inches(0.6), Inches(1.8), Inches(6.4), Inches(5), [
        f"Low is strictly dominant for both stations, so ({eq[0]}, {eq[1]}) is the unique pure Nash equilibrium.",
        f"At equilibrium each earns {rs(AN.matrix[eq][0])}; switching alone to High drops profit to "
        f"{rs(AN.matrix[(HIGH, LOW)][0])}.",
        f"High/High would give {rs(AN.matrix[(HIGH, HIGH)][0])} each ({rs(AN.joint[(HIGH, HIGH)])} combined) "
        f"but is unstable: undercutting earns {rs(AN.matrix[(LOW, HIGH)][0])}.",
        f"T > R > P > S: {rs(pd_a.temptation)} > {rs(pd_a.reward)} > {rs(pd_a.punishment)} > "
        f"{rs(pd_a.sucker)}, a Prisoner's Dilemma.",
        f"No interior mixed equilibrium: the indifference formula gives q* = {AN.mixed.q_b_first:.2f}, "
        "outside [0, 1].",
    ], 17)
    _picture(s, figs["joint"], Inches(7.2), Inches(1.8), w=Inches(5.8))
    _notes(s, f"Key viva point: the Nash equilibrium is stable against one-player deviations; the joint "
              f"maximum High/High is not. The firms forgo {rs(gap)} of combined profit, while drivers pay less.")

    # 8 Website
    s = _frame(prs, "Website and implementation", "An interactive website anyone can play", 8)
    _picture(s, SHOTS / "web_04_analysis.png", Inches(0.6), Inches(1.75), h=Inches(3.55))
    _picture(s, SHOTS / "web_03_simulator.png", Inches(5.4), Inches(1.75), h=Inches(3.55))
    _text(s, Inches(0.6), Inches(5.45), Inches(12.3), Inches(1.6), [
        "game_theory.py: model, best responses, Nash search, PD check, mixed strategies, sensitivity",
        "web/: Vite + GSAP site; model.js runs the same model in the browser  ·  49 pytest tests incl. JS parity",
        "Run locally:  cd web && npm install && npm run dev",
    ], 15, False, NAVY)
    _notes(s, "Live demo: in Play, set both stations to ₹12 and show each gains ₹200 by switching. "
              "In Analysis, every arrow points to (₹10, ₹10); press Present full screen. "
              "Open Change the market, set q to 0.70 and show two equilibria and the 50/50 mixed strategy. "
              "Then Restore baseline.")

    # 9 Results
    s = _frame(prs, "Results and sensitivity", "How the outcome changes with q and cost", 9)
    _picture(s, figs["q"], Inches(0.5), Inches(1.7), w=Inches(6.3))
    _picture(s, figs["heat"], Inches(6.9), Inches(1.7), w=Inches(6.1))
    _text(s, Inches(0.6), Inches(5.3), Inches(12.3), Inches(1.8), [
        f"q > {hi:.2f}: Low dominant, (Low, Low) unique, Prisoner's Dilemma.   "
        f"{lo:.3f} < q < {hi:.2f}: two pure equilibria.   q < {lo:.3f}: (High, High).",
        "Applications: pricing strategy for charging operators, understanding price wars, individual vs joint profit.",
        "Limitation: illustrative parameters, fixed demand, two price levels; not an empirical prediction.",
    ], 15, False, NAVY)
    _notes(s, "Undercutting only pays when it wins enough extra customers. The heatmap shows the same idea "
              "when costs change: lower costs make the low price's margin relatively larger.")

    # 10 Conclusion
    s = _frame(prs, "Conclusion and future scope", "What the model shows", 10)
    _bullets(s, Inches(0.6), Inches(1.8), Inches(6.2), Inches(5), [
        "Under the baseline, competition drives both stations to the low price.",
        "The equilibrium is stable, yet not the joint-profit maximum: a Prisoner's Dilemma.",
        "Whether undercutting pays depends on driver price sensitivity (q) and on margins.",
        "All results are computed, tested and explorable on the website.",
    ], 18)
    _text(s, Inches(7.2), Inches(1.8), Inches(5.6), Inches(0.5), "Future scope", 20, True, TEAL)
    _bullets(s, Inches(7.2), Inches(2.4), Inches(5.6), Inches(4.5), [
        "Continuous prices and elastic demand",
        "Charging capacity and queues",
        "Location-based (Hotelling) competition",
        "Repeated game: can High/High be sustained?",
        "More stations and real tariff data",
    ], 17)
    _rect(s, Inches(0.6), Inches(5.55), Inches(12.2), Inches(1.15), "#eef5f7", MSO_SHAPE.ROUNDED_RECTANGLE, TEAL)
    _text(s, Inches(0.9), Inches(5.7), Inches(11.7), Inches(0.9),
          f"Answer to the research question: the stations choose ({eq[0]}, {eq[1]}) and earn "
          f"{rs(AN.joint[eq])} together, while (High, High) would earn {rs(AN.joint[AN.joint_max[0]])}. "
          "Individual incentives and joint profit point in different directions.", 17, True, NAVY)
    _notes(s, "Restate the research question and the answer, then invite questions.")

    path = ROOT / "presentation.pptx"
    prs.save(path)
    return path


if __name__ == "__main__":
    FIG.mkdir(parents=True, exist_ok=True)
    figs = {"profits": fig_profits(), "joint": fig_joint(), "q": fig_q_sensitivity(),
            "heat": fig_heatmap(), "matrix": fig_matrix()}
    for p in figs.values():
        print("wrote", p.relative_to(ROOT))
    print("wrote", build_presentation(figs).relative_to(ROOT))
